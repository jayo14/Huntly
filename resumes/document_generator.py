"""Pure Python native ATS-friendly DOCX and PDF document generator for Huntly.

Generates validated, standards-compliant Office Open XML (.docx) and PDF (.pdf) documents
without subprocess dependencies, browser automation, or external binaries.
"""

import io
import re
from typing import Dict, Any, Optional

import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, HRFlowable, Table, TableStyle, KeepTogether
)


class DocumentGenerationError(Exception):
    """Raised when document generation fails."""
    pass


def _hex_to_rgb(hex_code: str) -> RGBColor:
    hex_clean = hex_code.lstrip('#')
    if len(hex_clean) == 6:
        r = int(hex_clean[0:2], 16)
        g = int(hex_clean[2:4], 16)
        b = int(hex_clean[4:6], 16)
        return RGBColor(r, g, b)
    return RGBColor(79, 70, 229)  # Default indigo


def _clean_text(val: Any) -> str:
    if val is None:
        return ""
    text = str(val).strip()
    # Replace non-printable ASCII/control chars that can corrupt XML
    text = re.sub(r'[\x00-\x08\x0B\x0C\x0E-\x1F\x7F]', '', text)
    return text


def generate_docx(data: Dict[str, Any]) -> bytes:
    """Generate an ATS-friendly, clean DOCX resume from structured data.

    Returns bytes starting with PK (Office Open XML ZIP format).
    """
    try:
        doc = Document()

        # Set 0.75 inch margins
        for section in doc.sections:
            section.top_margin = Inches(0.75)
            section.bottom_margin = Inches(0.75)
            section.left_margin = Inches(0.75)
            section.right_margin = Inches(0.75)

        accent_hex = data.get('accent_color', '#102a1f')
        accent_rgb = _hex_to_rgb(accent_hex)

        # 1. Header: Name & Tagline
        full_name = _clean_text(data.get('full_name', 'Candidate Profile'))
        name_p = doc.add_paragraph()
        name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        name_run = name_p.add_run(full_name)
        name_run.font.name = 'Calibri'
        name_run.font.size = Pt(22)
        name_run.font.bold = True
        name_run.font.color.rgb = accent_rgb
        name_p.paragraph_format.space_after = Pt(2)

        tagline = _clean_text(data.get('tagline', ''))
        if tagline:
            tag_p = doc.add_paragraph()
            tag_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            tag_run = tag_p.add_run(tagline)
            tag_run.font.name = 'Calibri'
            tag_run.font.size = Pt(11)
            tag_run.font.bold = True
            tag_run.font.color.rgb = RGBColor(80, 80, 80)
            tag_p.paragraph_format.space_after = Pt(4)

        # Contact Info Row
        contacts = []
        for key in ['email', 'phone', 'location', 'linkedin', 'github', 'portfolio']:
            val = _clean_text(data.get(key, ''))
            if val:
                contacts.append(val)

        if contacts:
            contact_p = doc.add_paragraph()
            contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            contact_run = contact_p.add_run("  •  ".join(contacts))
            contact_run.font.name = 'Calibri'
            contact_run.font.size = Pt(9.5)
            contact_run.font.color.rgb = RGBColor(100, 100, 100)
            contact_p.paragraph_format.space_after = Pt(14)

        def add_heading(title: str):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(title.upper())
            run.font.name = 'Calibri'
            run.font.size = Pt(11.5)
            run.font.bold = True
            run.font.color.rgb = accent_rgb

            # Add thin bottom border in XML
            pPr = p._p.get_or_add_pPr()
            pBdr = parse_xml(
                f'<w:pBdr {nsdecls("w")}>'
                f'<w:bottom w:val="single" w:sz="6" w:space="2" w:color="{accent_hex.lstrip("#")[:6]}"/>'
                f'</w:pBdr>'
            )
            pPr.append(pBdr)

        # 2. Professional Summary
        summary = _clean_text(data.get('summary', ''))
        if summary:
            add_heading("Professional Summary")
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(summary)
            run.font.name = 'Calibri'
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(40, 40, 40)

        # 3. Technical Skills
        skills = data.get('technical_skills', [])
        if skills:
            add_heading("Skills & Core Competencies")
            if isinstance(skills, list):
                for item in skills:
                    p = doc.add_paragraph()
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.left_indent = Inches(0.2)
                    if isinstance(item, dict):
                        category = _clean_text(item.get('category') or item.get('name') or 'Core')
                        skill_names = _clean_text(item.get('skills') or item.get('name') or '')
                        evidence = _clean_text(item.get('evidence', ''))
                        
                        r_cat = p.add_run(f"{category}: ")
                        r_cat.font.name = 'Calibri'
                        r_cat.font.bold = True
                        r_cat.font.size = Pt(9.5)
                        
                        r_skills = p.add_run(skill_names)
                        r_skills.font.name = 'Calibri'
                        r_skills.font.size = Pt(9.5)

                        if evidence:
                            r_ev = p.add_run(f" ({evidence})")
                            r_ev.font.name = 'Calibri'
                            r_ev.font.size = Pt(8.5)
                            r_ev.font.color.rgb = RGBColor(120, 120, 120)
                    else:
                        r = p.add_run(f"• {_clean_text(item)}")
                        r.font.name = 'Calibri'
                        r.font.size = Pt(9.5)

        # 4. Work Experience
        experiences = data.get('experiences', [])
        if experiences:
            add_heading("Work Experience")
            for exp in experiences:
                if not isinstance(exp, dict):
                    continue
                role = _clean_text(exp.get('role', ''))
                company = _clean_text(exp.get('company', ''))
                dates = _clean_text(exp.get('dates', ''))
                loc = _clean_text(exp.get('location', ''))

                header_p = doc.add_paragraph()
                header_p.paragraph_format.space_before = Pt(8)
                header_p.paragraph_format.space_after = Pt(2)
                header_p.paragraph_format.keep_with_next = True

                r_role = header_p.add_run(role)
                r_role.font.name = 'Calibri'
                r_role.font.bold = True
                r_role.font.size = Pt(10.5)

                if company:
                    r_comp = header_p.add_run(f"  |  {company}")
                    r_comp.font.name = 'Calibri'
                    r_comp.font.bold = True
                    r_comp.font.size = Pt(10)
                    r_comp.font.color.rgb = accent_rgb

                sub_parts = [p for p in [dates, loc] if p]
                if sub_parts:
                    r_sub = header_p.add_run(f"  ({', '.join(sub_parts)})")
                    r_sub.font.name = 'Calibri'
                    r_sub.font.italic = True
                    r_sub.font.size = Pt(9)
                    r_sub.font.color.rgb = RGBColor(100, 100, 100)

                # Bullets (Accomplishments / Responsibilities)
                bullets = exp.get('accomplishments') or exp.get('bullets') or []
                if isinstance(bullets, str):
                    bullets = [b.strip() for b in bullets.split('\n') if b.strip()]

                for b in bullets:
                    b_clean = _clean_text(b).lstrip('•-* ')
                    if not b_clean:
                        continue
                    bp = doc.add_paragraph()
                    bp.paragraph_format.left_indent = Inches(0.25)
                    bp.paragraph_format.space_before = Pt(1)
                    bp.paragraph_format.space_after = Pt(2)
                    r_b = bp.add_run(f"•  {b_clean}")
                    r_b.font.name = 'Calibri'
                    r_b.font.size = Pt(9.5)
                    r_b.font.color.rgb = RGBColor(40, 40, 40)

        # 5. Projects
        projects = data.get('projects', [])
        if projects:
            add_heading("Key Projects")
            for proj in projects:
                if not isinstance(proj, dict):
                    continue
                name = _clean_text(proj.get('name', ''))
                role = _clean_text(proj.get('role', ''))
                link = _clean_text(proj.get('link', ''))
                contribution = _clean_text(proj.get('contribution', ''))
                outcomes = _clean_text(proj.get('outcomes', ''))

                pp = doc.add_paragraph()
                pp.paragraph_format.space_before = Pt(6)
                pp.paragraph_format.space_after = Pt(2)
                pp.paragraph_format.keep_with_next = True

                r_pname = pp.add_run(name)
                r_pname.font.name = 'Calibri'
                r_pname.font.bold = True
                r_pname.font.size = Pt(10)

                if role:
                    r_prole = pp.add_run(f"  ({role})")
                    r_prole.font.name = 'Calibri'
                    r_prole.font.size = Pt(9)
                    r_prole.font.italic = True

                if link:
                    r_plink = pp.add_run(f"  [{link}]")
                    r_plink.font.name = 'Calibri'
                    r_plink.font.size = Pt(8.5)
                    r_plink.font.color.rgb = RGBColor(100, 100, 100)

                p_bullets = proj.get('bullets', [])
                if contribution:
                    p_bullets = [contribution] + (p_bullets if isinstance(p_bullets, list) else [])
                if outcomes:
                    p_bullets = (p_bullets if isinstance(p_bullets, list) else []) + [f"Outcome: {outcomes}"]

                for b in p_bullets:
                    b_clean = _clean_text(b).lstrip('•-* ')
                    if not b_clean:
                        continue
                    bp = doc.add_paragraph()
                    bp.paragraph_format.left_indent = Inches(0.25)
                    bp.paragraph_format.space_after = Pt(2)
                    r_b = bp.add_run(f"•  {b_clean}")
                    r_b.font.name = 'Calibri'
                    r_b.font.size = Pt(9.5)

        # 6. Education & Certifications
        education = data.get('education', [])
        certs = data.get('certifications', [])
        if education or certs:
            add_heading("Education & Certifications")
            for edu in education:
                if not isinstance(edu, dict):
                    continue
                deg = _clean_text(edu.get('degree', ''))
                inst = _clean_text(edu.get('institution', ''))
                dates = _clean_text(edu.get('dates', ''))
                ep = doc.add_paragraph()
                ep.paragraph_format.space_after = Pt(2)
                r_deg = ep.add_run(f"{deg} - {inst}")
                r_deg.font.name = 'Calibri'
                r_deg.font.bold = True
                r_deg.font.size = Pt(9.5)
                if dates:
                    r_dt = ep.add_run(f" ({dates})")
                    r_dt.font.name = 'Calibri'
                    r_dt.font.size = Pt(8.5)
                    r_dt.font.color.rgb = RGBColor(100, 100, 100)

            for c in certs:
                if not isinstance(c, dict):
                    continue
                cname = _clean_text(c.get('name', ''))
                iss = _clean_text(c.get('issuer', ''))
                dt = _clean_text(c.get('date', ''))
                cp = doc.add_paragraph()
                cp.paragraph_format.space_after = Pt(2)
                r_c = cp.add_run(f"• {cname}")
                r_c.font.name = 'Calibri'
                r_c.font.size = Pt(9.5)
                if iss or dt:
                    details = " - ".join([p for p in [iss, dt] if p])
                    r_iss = cp.add_run(f" ({details})")
                    r_iss.font.name = 'Calibri'
                    r_iss.font.size = Pt(8.5)
                    r_iss.font.color.rgb = RGBColor(100, 100, 100)

        # Save to buffer
        buffer = io.BytesIO()
        doc.save(buffer)
        output = buffer.getvalue()

        if not output or not output.startswith(b'PK'):
            raise DocumentGenerationError("Generated DOCX output does not have a valid Office Open XML header.")

        return output
    except Exception as e:
        if isinstance(e, DocumentGenerationError):
            raise
        raise DocumentGenerationError(f"DOCX generation failed: {e}") from e


def generate_pdf(data: Dict[str, Any]) -> bytes:
    """Generate an ATS-friendly, clean PDF resume from structured data.

    Returns bytes starting with %PDF-.
    """
    try:
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            leftMargin=40,
            rightMargin=40,
            topMargin=40,
            bottomMargin=40
        )

        accent_hex = data.get('accent_color', '#102a1f')
        accent_color = colors.HexColor(accent_hex)
        body_color = colors.HexColor('#222222')
        muted_color = colors.HexColor('#555555')

        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            'ResumeName',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=20,
            leading=24,
            textColor=accent_color,
            alignment=1, # Center
            spaceAfter=3,
        )

        tagline_style = ParagraphStyle(
            'ResumeTagline',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=10.5,
            leading=13,
            textColor=muted_color,
            alignment=1,
            spaceAfter=4,
        )

        contact_style = ParagraphStyle(
            'ResumeContact',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=8.5,
            leading=11,
            textColor=muted_color,
            alignment=1,
            spaceAfter=12,
        )

        heading_style = ParagraphStyle(
            'ResumeSectionHeading',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=11,
            leading=14,
            textColor=accent_color,
            spaceBefore=10,
            spaceAfter=4,
            keepWithNext=True,
        )

        body_style = ParagraphStyle(
            'ResumeBody',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9,
            leading=12.5,
            textColor=body_color,
            spaceAfter=4,
        )

        bullet_style = ParagraphStyle(
            'ResumeBullet',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=8.5,
            leading=11.5,
            textColor=body_color,
            leftIndent=14,
            spaceAfter=2,
        )

        story = []

        # 1. Header
        full_name = _clean_text(data.get('full_name', 'Candidate Profile'))
        story.append(Paragraph(full_name, title_style))

        tagline = _clean_text(data.get('tagline', ''))
        if tagline:
            story.append(Paragraph(tagline, tagline_style))

        contacts = []
        for key in ['email', 'phone', 'location', 'linkedin', 'github', 'portfolio']:
            val = _clean_text(data.get(key, ''))
            if val:
                contacts.append(val)
        if contacts:
            story.append(Paragraph(" &nbsp;•&nbsp; ".join(contacts), contact_style))

        def add_pdf_heading(title: str):
            story.append(Paragraph(title.upper(), heading_style))
            story.append(HRFlowable(width="100%", thickness=1, color=accent_color, spaceBefore=1, spaceAfter=5))

        # 2. Professional Summary
        summary = _clean_text(data.get('summary', ''))
        if summary:
            add_pdf_heading("Professional Summary")
            story.append(Paragraph(summary, body_style))

        # 3. Skills
        skills = data.get('technical_skills', [])
        if skills:
            add_pdf_heading("Skills & Core Competencies")
            for item in skills:
                if isinstance(item, dict):
                    category = _clean_text(item.get('category') or item.get('name') or 'Core')
                    skill_names = _clean_text(item.get('skills') or item.get('name') or '')
                    ev = _clean_text(item.get('evidence', ''))
                    ev_text = f" <font color='#777777'>({ev})</font>" if ev else ""
                    text = f"<b>{category}:</b> {skill_names}{ev_text}"
                    story.append(Paragraph(text, body_style))
                else:
                    story.append(Paragraph(f"• {_clean_text(item)}", bullet_style))

        # 4. Work Experience
        experiences = data.get('experiences', [])
        if experiences:
            add_pdf_heading("Work Experience")
            for exp in experiences:
                if not isinstance(exp, dict):
                    continue
                role = _clean_text(exp.get('role', ''))
                company = _clean_text(exp.get('company', ''))
                dates = _clean_text(exp.get('dates', ''))
                loc = _clean_text(exp.get('location', ''))

                header_parts = [f"<b>{role}</b>"]
                if company:
                    header_parts.append(f"<font color='{accent_hex}'><b>{company}</b></font>")
                sub = ", ".join([p for p in [dates, loc] if p])
                if sub:
                    header_parts.append(f"<font color='#666666'><i>({sub})</i></font>")

                story.append(Paragraph(" &nbsp;|&nbsp; ".join(header_parts), body_style))

                bullets = exp.get('accomplishments') or exp.get('bullets') or []
                if isinstance(bullets, str):
                    bullets = [b.strip() for b in bullets.split('\n') if b.strip()]

                for b in bullets:
                    b_clean = _clean_text(b).lstrip('•-* ')
                    if b_clean:
                        story.append(Paragraph(f"&bull; &nbsp;{b_clean}", bullet_style))
                story.append(Spacer(1, 4))

        # 5. Projects
        projects = data.get('projects', [])
        if projects:
            add_pdf_heading("Key Projects")
            for proj in projects:
                if not isinstance(proj, dict):
                    continue
                name = _clean_text(proj.get('name', ''))
                role = _clean_text(proj.get('role', ''))
                link = _clean_text(proj.get('link', ''))
                contrib = _clean_text(proj.get('contribution', ''))
                outcomes = _clean_text(proj.get('outcomes', ''))

                p_header = f"<b>{name}</b>"
                if role:
                    p_header += f" <i>({role})</i>"
                if link:
                    p_header += f" &nbsp; <font color='#666666'>[{link}]</font>"
                story.append(Paragraph(p_header, body_style))

                p_bullets = proj.get('bullets', [])
                if contrib:
                    p_bullets = [contrib] + (p_bullets if isinstance(p_bullets, list) else [])
                if outcomes:
                    p_bullets = (p_bullets if isinstance(p_bullets, list) else []) + [f"Outcome: {outcomes}"]

                for b in p_bullets:
                    b_clean = _clean_text(b).lstrip('•-* ')
                    if b_clean:
                        story.append(Paragraph(f"&bull; &nbsp;{b_clean}", bullet_style))
                story.append(Spacer(1, 3))

        # 6. Education & Certifications
        education = data.get('education', [])
        certs = data.get('certifications', [])
        if education or certs:
            add_pdf_heading("Education & Certifications")
            for edu in education:
                if not isinstance(edu, dict):
                    continue
                deg = _clean_text(edu.get('degree', ''))
                inst = _clean_text(edu.get('institution', ''))
                dates = _clean_text(edu.get('dates', ''))
                txt = f"<b>{deg}</b> &mdash; {inst}"
                if dates:
                    txt += f" <font color='#666666'>({dates})</font>"
                story.append(Paragraph(txt, body_style))

            for c in certs:
                if not isinstance(c, dict):
                    continue
                cname = _clean_text(c.get('name', ''))
                iss = _clean_text(c.get('issuer', ''))
                dt = _clean_text(c.get('date', ''))
                txt = f"&bull; <b>{cname}</b>"
                if iss or dt:
                    txt += f" <font color='#666666'>({(' - '.join([p for p in [iss, dt] if p]))})</font>"
                story.append(Paragraph(txt, body_style))

        doc.build(story)
        pdf_bytes = buffer.getvalue()

        if not pdf_bytes or not pdf_bytes.startswith(b'%PDF-'):
            raise DocumentGenerationError("Generated PDF output does not have a valid %PDF- header.")

        return pdf_bytes
    except Exception as e:
        if isinstance(e, DocumentGenerationError):
            raise
        raise DocumentGenerationError(f"PDF generation failed: {e}") from e


def generate_cover_letter_docx(candidate_name: str, candidate_contact: str, company_name: str, role_title: str, letter_body: str) -> bytes:
    """Generate a clean, professional cover letter DOCX."""
    data = {
        'full_name': candidate_name,
        'tagline': f"Application for {role_title} at {company_name}",
        'summary': letter_body,
        'email': candidate_contact,
    }
    return generate_docx(data)


def generate_cover_letter_pdf(candidate_name: str, candidate_contact: str, company_name: str, role_title: str, letter_body: str) -> bytes:
    """Generate a clean, professional cover letter PDF."""
    data = {
        'full_name': candidate_name,
        'tagline': f"Application for {role_title} at {company_name}",
        'summary': letter_body,
        'email': candidate_contact,
    }
    return generate_pdf(data)
